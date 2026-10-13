import csv
import hashlib
import json
import os
import re
import tempfile
import threading
import shutil
import uuid
from datetime import datetime
from io import BytesIO, StringIO

from flask import Flask, Response, jsonify, redirect, render_template, request, url_for, session

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "super-secret-cave-key")

DATA_DIR = "data"
if os.environ.get("VERCEL") or os.environ.get("AWS_LAMBDA_FUNCTION_NAME"):
    DATA_DIR = os.path.join(tempfile.gettempdir(), "data")
    if not os.path.exists(DATA_DIR):
        os.makedirs(DATA_DIR, exist_ok=True)
        bundled_data = os.path.join(os.path.dirname(__file__), "data")
        if os.path.exists(bundled_data):
            for item in os.listdir(bundled_data):
                src = os.path.join(bundled_data, item)
                dst = os.path.join(DATA_DIR, item)
                if os.path.isfile(src) and not os.path.exists(dst):
                    try:
                        shutil.copy2(src, dst)
                    except Exception:
                        pass
DIVISIONS = ("Division A", "Division B")
FILES = {
    "A": os.path.join(DATA_DIR, "transactions_gl.json"),
    "B": os.path.join(DATA_DIR, "transactions_bank.json"),
    "Q": os.path.join(DATA_DIR, "queue.json"),
    "P": os.path.join(DATA_DIR, "processed.json"),
    "I": os.path.join(DATA_DIR, "imports.json"),
    "M": os.path.join(DATA_DIR, "mapping_presets.json"),
    "C": os.path.join(DATA_DIR, "current_division.json"),
}
TABLE_LABELS = {"A": "General Ledger", "B": "Bank Statement"}
SOURCE_LABELS = {"A": "GL", "B": "BANK"}
TABLE_COLUMN_ORDER = [
    "date",
    "reference",
    "document_no",
    "counterparty",
    "currency",
    "currency_amount",
    "balance",
    "notes",
]

# DIVISIONS = {
#   "Division A",
#   "Division B"
# }
FIELD_LABELS = {
    "description": "Description",
    "amount": "Amount",
    "date": "Date",
    "reference": "Reference",
    "document_no": "Document No",
    "counterparty": "Counterparty",
    "currency": "Currency",
    "currency_amount": "Currency Amount",
    "balance": "Balance",
    "notes": "Notes",
}
CANONICAL_FIELDS = [
    "ignore",
    "description",
    "amount",
    "currency",
    "currency_amount",
    "debit",
    "credit",
    "date",
    "reference",
    "balance",
    "document_no",
    "counterparty",
    "notes",
]
PREVIEW_LIMIT = 8
FILE_LOCK = threading.Lock()


def normalize_division(division):
    clean = str(division or "").strip().upper()
    for d in DIVISIONS:
        if d.upper() == clean:
            return d
    return "Division A"


def utc_now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def atomic_write_json(path, data):
    try:
        dir_name = os.path.dirname(path) or "."
        os.makedirs(dir_name, exist_ok=True)
        with tempfile.NamedTemporaryFile("w", dir=dir_name, delete=False) as tmp:
            json.dump(data, tmp, indent=4, ensure_ascii=True)
            tmp.flush()
            os.fsync(tmp.fileno())
            temp_name = tmp.name
        os.replace(temp_name, path)
    except (OSError, PermissionError):
        try:
            fallback = os.path.join(tempfile.gettempdir(), os.path.basename(path))
            with open(fallback, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=4, ensure_ascii=True)
        except Exception:
            pass


def get_storage_path(key, division=None):
    if key in ("A", "B", "Q", "P", "I"):
        resolved_division = normalize_division(division or get_active_division())
        if resolved_division == "Division A":
            return FILES[key]
        base_name = os.path.basename(FILES[key])
        stem, ext = os.path.splitext(base_name)
        return os.path.join(DATA_DIR, f"{stem}_{resolved_division.lower().replace(' ', '_')}{ext}")
    return FILES[key]


def init_db():
    try:
        os.makedirs(DATA_DIR, exist_ok=True)
        for key in ("A", "B", "Q", "P", "I"):
            for division in DIVISIONS:
                path = get_storage_path(key, division)
                if not os.path.exists(path):
                    atomic_write_json(path, [])
        for key in ("M",):
            if not os.path.exists(FILES[key]):
                atomic_write_json(FILES[key], [])
        if not os.path.exists(FILES["C"]):
            atomic_write_json(FILES["C"], {"active_division": "Division A"})
        migrate_legacy_data()
    except (OSError, PermissionError):
        pass


def load_json(key, default=None, division=None):
    path = get_storage_path(key, division)
    if default is None:
        default = []
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        try:
            bundled = os.path.join(os.path.dirname(__file__), "data", os.path.basename(path))
            if os.path.exists(bundled):
                with open(bundled, "r", encoding="utf-8") as f:
                    return json.load(f)
        except Exception:
            pass
        return default


def save_json(key, data, division=None):
    with FILE_LOCK:
        atomic_write_json(get_storage_path(key, division), data)


def get_active_division():
    try:
        if "active_division" in session:
            return normalize_division(session["active_division"])
    except Exception:
        pass
    payload = load_json("C", {"active_division": "Division A"})
    if not isinstance(payload, dict):
        payload = {"active_division": "Division A"}
    active = normalize_division(payload.get("active_division"))
    return active


def set_active_division(division):
    active = normalize_division(division)
    try:
        session["active_division"] = active
    except Exception:
        pass
    try:
        save_json("C", {"active_division": active})
    except Exception:
        pass
    return active


def migrate_legacy_data():
    legacy_files = {
        "A": os.path.join(DATA_DIR, "table_a.json"),
        "B": os.path.join(DATA_DIR, "table_b.json"),
    }
    for key in ("A", "B"):
        current = load_json(key, [], division="Division A")
        legacy_path = legacy_files[key]
        if current or not os.path.exists(legacy_path):
            continue
        try:
            with open(legacy_path, "r", encoding="utf-8") as f:
                legacy_rows = json.load(f)
        except (FileNotFoundError, json.JSONDecodeError):
            legacy_rows = []
        if not legacy_rows:
            continue
        migrated = []
        for row in legacy_rows:
            migrated.append(
                build_transaction(
                    source_type=key,
                    import_batch_id="legacy-migration",
                    file_name=os.path.basename(legacy_path),
                    sheet_name="legacy",
                    mapped_fields={"description": "item", "amount": "amount"},
                    raw_row=row,
                    description=row.get("description")
                    or row.get("item")
                    or "Unknown Entry",
                    amount=parse_amount(row.get("amount")),
                    date=row.get("date", ""),
                    reference=row.get("reference", ""),
                    balance=parse_amount(row.get("balance")),
                    document_no=row.get("document_no", ""),
                    counterparty=row.get("counterparty", ""),
                    notes=row.get("notes", ""),
                    existing_id=row.get("id"),
                )
            )
        save_json(key, migrated, division="Division A")


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_amount(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)

    raw = clean_text(value)
    if not raw:
        return None

    negative = False
    if raw.startswith("(") and raw.endswith(")"):
        negative = True
        raw = raw[1:-1]

    normalized = (
        raw.replace("$", "")
        .replace(",", "")
        .replace("CR", "")
        .replace("DR", "")
        .strip()
    )
    if normalized.endswith("-"):
        negative = True
        normalized = normalized[:-1]

    try:
        amount = float(normalized)
    except ValueError:
        return None
    return -amount if negative else amount


def normalize_date(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return clean_text(value)


def display_date(value):
    raw = clean_text(value)
    if not raw:
        return ""

    patterns = [
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%m-%d-%Y",
        "%m/%d/%Y",
        "%d-%b-%Y",
        "%d/%b/%Y",
        "%d-%B-%Y",
        "%d/%B/%Y",
    ]
    for pattern in patterns:
        try:
            parsed = datetime.strptime(raw, pattern)
            return parsed.strftime("%d/%m/%Y")
        except ValueError:
            continue
    return raw


def amount_or_zero(value):
    parsed = parse_amount(value)
    return parsed if parsed is not None else 0.0


def get_recon_amount(item):
    amount = amount_or_zero(item.get("amount"))
    return amount if item.get("source_type") == "A" else -amount


def calculate_variance(items):
    return sum(get_recon_amount(item) for item in items)


def calculate_side_total(items, source_type):
    if source_type == "A":
        return sum(
            get_recon_amount(item) for item in items if item.get("source_type") == "A"
        )
    return sum(
        get_recon_amount(item) for item in items if item.get("source_type") == "B"
    )


def get_export_side_amount(item, side_label):
    amount = amount_or_zero(item.get("amount"))
    return -abs(amount) if side_label == "BANK" else amount


def transaction_store_key(source_type):
    if source_type not in ("A", "B"):
        raise ValueError("Invalid source type")
    return source_type


def read_transactions(source_type, division=None):
    resolved_division = normalize_division(division or get_active_division())
    records = load_json(
        transaction_store_key(source_type), [], division=resolved_division
    )
    normalized = []
    changed = False
    for record in records:
        upgraded, upgraded_changed = upgrade_transaction_record(record, source_type)
        normalized.append(upgraded)
        changed = changed or upgraded_changed
    if changed:
        save_json(source_type, normalized, division=resolved_division)
    return normalized


def read_queue(division=None):
    resolved_division = normalize_division(division or get_active_division())
    records = load_json("Q", [], division=resolved_division)
    normalized = []
    changed = False
    for record in records:
        source_type = record.get("source_type") or record.get("source")
        upgraded, upgraded_changed = upgrade_transaction_record(record, source_type)
        normalized.append(upgraded)
        changed = changed or upgraded_changed
    if changed:
        save_json("Q", normalized, division=resolved_division)
    return normalized


def read_processed(division=None):
    resolved_division = normalize_division(division or get_active_division())
    groups = load_json("P", [], division=resolved_division)
    normalized_groups = []
    changed = False
    for group in groups:
        normalized_group = dict(group)
        items_a = []
        items_b = []
        for item in group.get("items_a", []):
            upgraded, upgraded_changed = upgrade_transaction_record(item, "A")
            items_a.append(upgraded)
            changed = changed or upgraded_changed
        for item in group.get("items_b", []):
            upgraded, upgraded_changed = upgrade_transaction_record(item, "B")
            items_b.append(upgraded)
            changed = changed or upgraded_changed
        normalized_group["items_a"] = items_a
        normalized_group["items_b"] = items_b
        recalculated_group_sum = calculate_variance(items_a + items_b)
        expected_group_sum = f"{recalculated_group_sum:.2f}"
        if normalized_group.get("group_sum") != expected_group_sum:
            normalized_group["group_sum"] = expected_group_sum
            changed = True
        normalized_groups.append(normalized_group)
    if changed:
        save_json("P", normalized_groups, division=resolved_division)
    return normalized_groups


def upgrade_transaction_record(record, default_source_type=None):
    upgraded = dict(record)
    changed = False

    source_type = (
        upgraded.get("source_type")
        or upgraded.get("source")
        or default_source_type
        or "A"
    )
    description = upgraded.get("description") or upgraded.get("item") or "Unknown Entry"
    amount = parse_amount(upgraded.get("amount"))
    if amount is None:
        amount = 0.0

    base = {
        "id": upgraded.get("id")
        or generate_stable_id(
            source_type,
            description,
            amount,
            upgraded.get("date", ""),
            upgraded.get("reference", ""),
            upgraded.get("document_no", ""),
        ),
        "source_type": source_type,
        "source": source_type,
        "item": description,
        "description": description,
        "amount": amount,
        "currency": clean_text(upgraded.get("currency", "")),
        "currency_amount": parse_amount(upgraded.get("currency_amount")),
        "date": normalize_date(upgraded.get("date", "")),
        "reference": clean_text(upgraded.get("reference", "")),
        "balance": parse_amount(upgraded.get("balance")),
        "document_no": clean_text(upgraded.get("document_no", "")),
        "counterparty": clean_text(upgraded.get("counterparty", "")),
        "notes": clean_text(upgraded.get("notes", "")),
        "raw_row": upgraded.get("raw_row", {}),
        "mapped_fields": upgraded.get("mapped_fields", {}),
        "import_batch_id": upgraded.get("import_batch_id", ""),
        "import_file_name": clean_text(upgraded.get("import_file_name", "")),
        "sheet_name": clean_text(upgraded.get("sheet_name", "")),
        "created_at": upgraded.get("created_at") or utc_now(),
    }

    if any(upgraded.get(k) != v for k, v in base.items()):
        changed = True
    upgraded.update(base)
    return upgraded, changed


def build_transaction(
    source_type,
    import_batch_id,
    file_name,
    sheet_name,
    mapped_fields,
    raw_row,
    description,
    amount,
    currency="",
    currency_amount=None,
    date="",
    reference="",
    balance=None,
    document_no="",
    counterparty="",
    notes="",
    existing_id=None,
):
    description = clean_text(description) or "Unknown Entry"
    amount = 0.0 if amount is None else float(amount)
    currency = clean_text(currency).upper()
    currency_amount = parse_amount(currency_amount)
    date = normalize_date(date)
    reference = clean_text(reference)
    document_no = clean_text(document_no)
    counterparty = clean_text(counterparty)
    notes = clean_text(notes)

    return {
        "id": existing_id
        or generate_stable_id(
            source_type, description, amount, date, reference, document_no
        ),
        "source_type": source_type,
        "source": source_type,
        "item": description,
        "description": description,
        "amount": amount,
        "currency": currency,
        "currency_amount": currency_amount,
        "date": date,
        "reference": reference,
        "balance": balance,
        "document_no": document_no,
        "counterparty": counterparty,
        "notes": notes,
        "raw_row": raw_row,
        "mapped_fields": mapped_fields,
        "import_batch_id": import_batch_id,
        "import_file_name": file_name,
        "sheet_name": sheet_name,
        "created_at": utc_now(),
    }


def generate_stable_id(
    source_type, description, amount, date="", reference="", document_no=""
):
    combined = "|".join(
        [
            clean_text(source_type).lower(),
            clean_text(description).lower(),
            f"{float(amount):.2f}",
            clean_text(date).lower(),
            clean_text(reference).lower(),
            clean_text(document_no).lower(),
        ]
    )
    return hashlib.md5(combined.encode("utf-8")).hexdigest()[:12]


def get_all_known_ids():
    ids = set()
    for key in ("A", "B"):
        for item in read_transactions(key):
            ids.add(str(item.get("id")))
    for item in read_queue():
        ids.add(str(item.get("id")))
    for group in read_processed():
        for item in group.get("items_a", []):
            ids.add(str(item.get("id")))
        for item in group.get("items_b", []):
            ids.add(str(item.get("id")))
    return ids


def transaction_meta_lines(item):
    parts = []
    if item.get("date"):
        parts.append(f"Date: {item['date']}")
    if item.get("currency_amount") not in (None, "") and item.get("currency"):
        parts.append(f"{item['currency']} {float(item['currency_amount']):.2f}")
    elif item.get("currency"):
        parts.append(f"Currency: {item['currency']}")
    if item.get("reference"):
        parts.append(f"Ref: {item['reference']}")
    if item.get("document_no"):
        parts.append(f"Doc: {item['document_no']}")
    if item.get("counterparty"):
        parts.append(f"Party: {item['counterparty']}")
    if item.get("balance") not in (None, ""):
        parts.append(f"Balance: {float(item['balance']):.2f}")
    return parts[:3]


def group_meta(item):
    lines = transaction_meta_lines(item)
    if item.get("notes"):
        lines.append(f"Notes: {item['notes']}")
    return lines[:4]


def format_field_value(value, field_name=None):
    if value in (None, ""):
        return "—"
    if field_name == "date":
        return display_date(value)
    if field_name in ("amount", "balance", "currency_amount"):
        parsed = parse_amount(value)
        if parsed is not None:
            return f"{parsed:.2f}"
    return str(value)


def get_available_table_columns(transactions):
    available = []
    for field in TABLE_COLUMN_ORDER:
        if any(
            clean_text(item.get(field)) or item.get(field) not in (None, "")
            for item in transactions
        ):
            available.append(
                {
                    "key": field,
                    "label": FIELD_LABELS.get(field, field.replace("_", " ").title()),
                }
            )
    return available


def get_default_visible_columns(transactions):
    available_keys = [
        column["key"] for column in get_available_table_columns(transactions)
    ]
    preferred = ["date", "reference", "currency", "currency_amount"]
    visible = [key for key in preferred if key in available_keys]
    return visible[:4]


def get_transaction_detail_rows(item):
    rows = []
    mapped_fields = item.get("mapped_fields") or {}
    selected_fields = []

    if mapped_fields:
        for field in ["description", "amount"] + TABLE_COLUMN_ORDER:
            if field in mapped_fields and field not in selected_fields:
                selected_fields.append(field)
    else:
        for field in ["description", "amount"] + TABLE_COLUMN_ORDER:
            value = item.get(field)
            if value not in (None, ""):
                selected_fields.append(field)

    for field in selected_fields:
        value = item.get(field)
        if value in (None, ""):
            continue
        label = FIELD_LABELS.get(field, field.replace("_", " ").title())
        rows.append({"label": label, "value": format_field_value(value, field)})

    rows.append(
        {
            "label": "Reconciliation Amount",
            "value": format_field_value(get_recon_amount(item), "amount"),
        }
    )
    return rows


def get_excel_display_amount(item):
    amount = parse_amount(item.get("amount"))
    if amount is None:
        return 0.0
    if item.get("source_type") == "B":
        return -abs(amount)
    return amount


def autosize_worksheet_columns(worksheet):
    for column_cells in worksheet.columns:
        values = [
            len(clean_text(cell.value))
            for cell in column_cells
            if cell.value not in (None, "")
        ]
        if not values:
            continue
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(values) + 2, 40
        )


def apply_excel_number_formats(worksheet):
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"


def build_processed_export_workbook(processed):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Processed Summary"

    summary_headers = [
        "Reco ID",
        "Timestamp",
        "Variance",
        "GL Count",
        "Bank Count",
        "GL Total",
        "Bank Total",
        "GL Reference(s)",
        "Adjustment Count",
        "Adjustment Total",
    ]
    summary_sheet.append(summary_headers)

    all_items = []
    raw_field_order = []
    seen_raw_fields = set()

    for group in processed:
        gl_total = sum(
            get_excel_display_amount(item) for item in group.get("items_a", [])
        )
        bank_total = sum(
            get_excel_display_amount(item) for item in group.get("items_b", [])
        )
        all_group_items = list(group.get("items_a", [])) + list(
            group.get("items_b", [])
        )
        adjustment_items = [
            item for item in all_group_items if is_adjustment_item(item)
        ]
        adjustment_total = sum(get_recon_amount(item) for item in adjustment_items)
        gl_references = ", ".join(
            clean_text(item.get("reference"))
            for item in group.get("items_a", [])
            if clean_text(item.get("reference"))
        )
        summary_sheet.append(
            [
                group.get("group_id", ""),
                group.get("timestamp", ""),
                calculate_variance(all_group_items),
                len(group.get("items_a", [])),
                len(group.get("items_b", [])),
                gl_total,
                bank_total,
                gl_references,
                len(adjustment_items),
                adjustment_total,
            ]
        )

        for side_key, side_label in (("items_a", "GL"), ("items_b", "BANK")):
            for item in group.get(side_key, []):
                raw_row = item.get("raw_row") or {}
                for raw_key in raw_row.keys():
                    clean_key = clean_text(raw_key)
                    if clean_key and clean_key not in seen_raw_fields:
                        seen_raw_fields.add(clean_key)
                        raw_field_order.append(clean_key)
                all_items.append((group, side_label, item))

    detail_sheet = workbook.create_sheet("Processed Lines")
    detail_headers = [
        "Reco ID",
        "Timestamp",
        "Side",
        "Entry Type",
        "Description",
        "Imported Amount",
        "Reconciliation Amount",
        "Date",
        "Reference",
        "Document No",
        "Counterparty",
        "Currency",
        "Currency Amount",
        "Balance",
        "Notes",
        "Import File",
        "Sheet Name",
    ] + raw_field_order
    detail_sheet.append(detail_headers)

    for group, side_label, item in all_items:
        row = [
            group.get("group_id", ""),
            group.get("timestamp", ""),
            side_label,
            "Adjustment" if is_adjustment_item(item) else "Imported Transaction",
            item.get("description", ""),
            get_excel_display_amount(item),
            get_recon_amount(item),
            display_date(item.get("date", "")),
            item.get("reference", ""),
            item.get("document_no", ""),
            item.get("counterparty", ""),
            item.get("currency", ""),
            parse_amount(item.get("currency_amount"))
            if item.get("currency_amount") not in (None, "")
            else "",
            parse_amount(item.get("balance"))
            if item.get("balance") not in (None, "")
            else "",
            item.get("notes", ""),
            item.get("import_file_name", ""),
            item.get("sheet_name", ""),
        ]
        raw_row = item.get("raw_row") or {}
        for raw_key in raw_field_order:
            row.append(raw_row.get(raw_key, ""))
        detail_sheet.append(row)

    for sheet in (summary_sheet, detail_sheet):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        apply_excel_number_formats(sheet)
        autosize_worksheet_columns(sheet)

    return workbook


def build_transactions_export_workbook(source_type, transactions):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{SOURCE_LABELS[source_type]} Transactions"

    dynamic_fields = [
        column["key"] for column in get_available_table_columns(transactions)
    ]
    headers = (
        ["Description"]
        + [FIELD_LABELS.get(field, field.title()) for field in dynamic_fields]
        + [
            "Amount",
            "Reconciliation Amount",
        ]
    )
    sheet.append(headers)

    for item in transactions:
        row = [item.get("description", "")]
        for field in dynamic_fields:
            value = item.get(field)
            if field == "date":
                row.append(display_date(value))
            elif field in ("amount", "balance", "currency_amount"):
                parsed = parse_amount(value)
                row.append(parsed if parsed is not None else "")
            else:
                row.append(clean_text(value))
        row.extend([amount_or_zero(item.get("amount")), get_recon_amount(item)])
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    apply_excel_number_formats(sheet)
    autosize_worksheet_columns(sheet)
    return workbook


def populate_transactions_sheet(sheet, transactions):
    dynamic_fields = [
        column["key"] for column in get_available_table_columns(transactions)
    ]
    headers = (
        ["Description"]
        + [FIELD_LABELS.get(field, field.title()) for field in dynamic_fields]
        + [
            "Amount",
            "Reconciliation Amount",
            "Import File",
            "Sheet Name",
        ]
    )
    sheet.append(headers)

    for item in transactions:
        row = [item.get("description", "")]
        for field in dynamic_fields:
            value = item.get(field)
            if field == "date":
                row.append(display_date(value))
            elif field in ("amount", "balance", "currency_amount"):
                parsed = parse_amount(value)
                row.append(parsed if parsed is not None else "")
            else:
                row.append(clean_text(value))
        row.extend(
            [
                amount_or_zero(item.get("amount")),
                get_recon_amount(item),
                item.get("import_file_name", ""),
                item.get("sheet_name", ""),
            ]
        )
        sheet.append(row)


def is_adjustment_item(item):
    return item.get("import_batch_id") == "manual-adjustment" or clean_text(
        item.get("description", "")
    ).startswith("ADJ:")


def parse_csv_bytes(file_bytes):
    text = file_bytes.decode("utf-8-sig")
    reader = csv.reader(StringIO(text))
    rows = list(reader)
    if not rows:
        raise ValueError("The uploaded CSV is empty.")
    headers = [clean_text(h) or f"Column {idx + 1}" for idx, h in enumerate(rows[0])]
    data_rows = []
    for row in rows[1:]:
        padded = row + [""] * (len(headers) - len(row))
        data_rows.append(
            {
                headers[idx]: clean_text(padded[idx]) if idx < len(padded) else ""
                for idx in range(len(headers))
            }
        )
    return [{"name": "Sheet1", "headers": headers, "rows": data_rows}]


def excel_date_to_display(value, number_format):
    if not isinstance(value, datetime):
        return clean_text(value)

    fmt = clean_text(number_format)
    if not fmt:
        return value.strftime("%Y-%m-%d")

    # Remove Excel-specific decorations like colors/locales/conditions.
    fmt = re.sub(r"\[[^\]]+\]", "", fmt)
    fmt = re.sub(r'"[^"]*"', "", fmt)
    lower_fmt = fmt.lower()

    token_map = {
        "yyyy": "%Y",
        "yy": "%y",
        "mmmm": "%B",
        "mmm": "%b",
        "mm": "%m",
        "m": "%m",
        "dd": "%d",
        "d": "%d",
    }

    parts = re.split(r"([a-zA-Z]+)", lower_fmt)
    converted_parts = []
    for part in parts:
        converted_parts.append(token_map.get(part, part))
    strftime_fmt = "".join(converted_parts)

    # If conversion failed badly, fall back to the familiar ISO-like format.
    if "%" not in strftime_fmt:
        return value.strftime("%Y-%m-%d")

    try:
        return value.strftime(strftime_fmt)
    except ValueError:
        return value.strftime("%Y-%m-%d")


def parse_xlsx_bytes(file_bytes):
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise ValueError(
            "Excel support requires the openpyxl package. Install dependencies first."
        ) from exc
    workbook = load_workbook(
        filename=BytesIO(file_bytes), data_only=True, read_only=False
    )
    sheets = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows())
        if not rows:
            sheets.append({"name": sheet.title, "headers": [], "rows": []})
            continue
        headers = [
            clean_text(cell.value) or f"Column {idx + 1}"
            for idx, cell in enumerate(rows[0])
        ]
        data_rows = []
        for row in rows[1:]:
            values = list(row or [])
            padded = values + [None] * (len(headers) - len(values))
            row_dict = {}
            for idx, header in enumerate(headers):
                cell = padded[idx] if idx < len(padded) else None
                value = cell.value if cell is not None else ""
                if cell is not None and cell.is_date and isinstance(value, datetime):
                    row_dict[header] = excel_date_to_display(value, cell.number_format)
                else:
                    row_dict[header] = clean_text(value)
            data_rows.append(row_dict)
        sheets.append({"name": sheet.title, "headers": headers, "rows": data_rows})
    return sheets


def load_uploaded_sheets(file_storage):
    filename = clean_text(file_storage.filename)
    if not filename:
        raise ValueError("Please choose a file to upload.")

    file_bytes = file_storage.read()
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        sheets = parse_csv_bytes(file_bytes)
    elif lower_name.endswith(".xlsx"):
        sheets = parse_xlsx_bytes(file_bytes)
    else:
        raise ValueError("Only .csv and .xlsx files are supported.")

    return filename, sheets


def detect_mapping(headers):
    alias_map = {
        "description": [
            "description",
            "gl description",
            "narration",
            "details",
            "item",
        ],
        "currency": ["currency", "currency code", "ccy", "curr"],
        "currency_amount": [
            "currency amount",
            "foreign amount",
            "original amount",
            "amount in currency",
            "transaction currency amount",
        ],
        "amount": [
            "amount in sar",
            "sar amount",
            "amount sar",
            "net amount",
            "value",
            "transaction amount",
            "amount",
        ],
        "debit": ["debit", "withdrawal", "payment"],
        "credit": ["credit", "deposit", "receipt"],
        "date": ["date", "posting date", "transaction date"],
        "reference": ["reference", "ref", "voucher", "journal ref"],
        "balance": ["balance", "closing balance", "gl balance", "running balance"],
        "document_no": ["document no", "document number", "doc no", "invoice no"],
        "counterparty": ["counterparty", "customer", "vendor", "party", "name"],
        "notes": ["notes", "memo", "remarks", "comment"],
    }
    guesses = {}
    lowered_headers = {header: header.strip().lower() for header in headers}
    for header, lowered in lowered_headers.items():
        guessed = "ignore"
        for canonical, aliases in alias_map.items():
            if lowered in aliases:
                guessed = canonical
                break
        if guessed == "ignore":
            for canonical, aliases in alias_map.items():
                if any(alias in lowered for alias in aliases):
                    guessed = canonical
                    break
        guesses[header] = guessed
    return guesses


def get_presets():
    presets = load_json("M", [])
    by_source = {"A": [], "B": []}
    for preset in presets:
        source_type = preset.get("table_type")
        if source_type in by_source:
            by_source[source_type].append(preset)
    return by_source


def save_preset(table_type, name, mapping):
    presets = load_json("M", [])
    preset_id = hashlib.md5(f"{table_type}:{name}".encode("utf-8")).hexdigest()[:12]
    payload = {
        "id": preset_id,
        "table_type": table_type,
        "name": clean_text(name),
        "mapping": mapping,
        "updated_at": utc_now(),
    }
    updated = [p for p in presets if p.get("id") != preset_id]
    updated.insert(0, payload)
    save_json("M", updated)
    return payload


def build_import_response(import_entry):
    active_sheet_name = import_entry.get("active_sheet")
    active_sheet = next(
        (
            sheet
            for sheet in import_entry.get("sheets", [])
            if sheet.get("name") == active_sheet_name
        ),
        {"name": "", "headers": [], "rows": []},
    )
    detected_mapping = detect_mapping(active_sheet.get("headers", []))
    return {
        "import_id": import_entry["id"],
        "table_type": import_entry["table_type"],
        "file_name": import_entry["file_name"],
        "sheet_names": [sheet["name"] for sheet in import_entry.get("sheets", [])],
        "active_sheet": active_sheet_name,
        "headers": active_sheet.get("headers", []),
        "preview_rows": active_sheet.get("rows", [])[:PREVIEW_LIMIT],
        "row_count": len(active_sheet.get("rows", [])),
        "detected_mapping": detected_mapping,
        "canonical_fields": CANONICAL_FIELDS,
    }


def validate_mapping(headers, mapping):
    assigned = {
        field: [header for header, target in mapping.items() if target == field]
        for field in CANONICAL_FIELDS
    }
    errors = []
    header_counts = {}
    for header in headers:
        header_counts[header] = header_counts.get(header, 0) + 1
    duplicate_headers = [header for header, count in header_counts.items() if count > 1]
    if duplicate_headers:
        errors.append(f"Duplicate headers detected: {', '.join(duplicate_headers)}")
    if not assigned["description"]:
        errors.append("Map at least one column to description.")
    has_amount = bool(assigned["amount"])
    has_split_amount = bool(assigned["debit"] or assigned["credit"])
    if not has_amount and not has_split_amount:
        errors.append("Map an amount column or at least one debit/credit column.")
    return errors


def normalize_row(row, table_type, mapping, import_entry, active_sheet_name):
    field_values = {field: "" for field in CANONICAL_FIELDS}
    mapped_fields = {}
    for header, canonical_field in mapping.items():
        if canonical_field not in CANONICAL_FIELDS or canonical_field == "ignore":
            continue
        field_values[canonical_field] = row.get(header, "")
        mapped_fields[canonical_field] = header

    description = clean_text(field_values["description"])
    amount = parse_amount(field_values["amount"])
    if amount is None:
        debit = amount_or_zero(field_values["debit"])
        credit = amount_or_zero(field_values["credit"])
        amount = credit - debit
        if (
            clean_text(field_values["debit"]) == ""
            and clean_text(field_values["credit"]) == ""
        ):
            amount = None

    if not description:
        return None, "Missing description"
    if amount is None:
        return None, "Missing amount"

    balance = parse_amount(field_values["balance"])
    transaction = build_transaction(
        source_type=table_type,
        import_batch_id=import_entry["id"],
        file_name=import_entry["file_name"],
        sheet_name=active_sheet_name,
        mapped_fields=mapped_fields,
        raw_row=row,
        description=description,
        amount=amount,
        date=field_values["date"],
        currency=field_values["currency"],
        currency_amount=field_values["currency_amount"],
        reference=field_values["reference"],
        balance=balance,
        document_no=field_values["document_no"],
        counterparty=field_values["counterparty"],
        notes=field_values["notes"],
    )
    return transaction, None


def update_import_record(import_id, updater):
    active_division = get_active_division()
    imports = load_json("I", [], division=active_division)
    updated_imports = []
    updated_entry = None
    for entry in imports:
        if entry.get("id") == import_id:
            updated_entry = updater(dict(entry))
            updated_imports.append(updated_entry)
        else:
            updated_imports.append(entry)
    if updated_entry is None:
        return None
    save_json("I", updated_imports, division=active_division)
    return updated_entry


@app.route("/switch_division/<division>", methods=["POST"])
def switch_division(division):
    set_active_division(division)
    return redirect(url_for("index"))


@app.route("/")
def index():
    active_division = get_active_division()
    table_a = read_transactions("A", active_division)
    table_b = read_transactions("B", active_division)
    queue = read_queue(active_division)
    processed = read_processed(active_division)
    presets = get_presets()

    q_total = calculate_variance(queue)

    return render_template(
        "index.html",
        divisions=DIVISIONS,
        active_division=active_division,
        table_a=table_a,
        table_b=table_b,
        table_a_columns=get_available_table_columns(table_a),
        table_b_columns=get_available_table_columns(table_b),
        table_a_default_columns=get_default_visible_columns(table_a),
        table_b_default_columns=get_default_visible_columns(table_b),
        queue=queue,
        processed=processed,
        queue_total=f"{q_total:.2f}",
        presets=presets,
        canonical_fields=CANONICAL_FIELDS,
        transaction_meta_lines=transaction_meta_lines,
        group_meta=group_meta,
        format_field_value=format_field_value,
        get_transaction_detail_rows=get_transaction_detail_rows,
    )


@app.route("/upload_import/<table_type>", methods=["POST"])
def upload_import(table_type):
    if table_type not in ("A", "B"):
        return jsonify({"error": "Invalid table type"}), 400
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    file = request.files["file"]
    try:
        file_name, sheets = load_uploaded_sheets(file)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": f"Could not read file: {exc}"}), 500

    if not sheets:
        return jsonify({"error": "No worksheets found in the uploaded file."}), 400

    import_entry = {
        "id": str(uuid.uuid4())[:8],
        "table_type": table_type,
        "file_name": file_name,
        "created_at": utc_now(),
        "active_sheet": sheets[0]["name"],
        "sheets": sheets,
    }

    active_division = get_active_division()
    imports = load_json("I", [], division=active_division)
    imports = [entry for entry in imports if entry.get("table_type") != table_type][:20]
    imports.insert(0, import_entry)
    save_json("I", imports, division=active_division)
    return jsonify(build_import_response(import_entry))


@app.route("/set_import_sheet/<import_id>", methods=["POST"])
def set_import_sheet(import_id):
    sheet_name = clean_text(
        request.json.get("sheet_name")
        if request.is_json
        else request.form.get("sheet_name")
    )
    if not sheet_name:
        return jsonify({"error": "Sheet name is required"}), 400

    def updater(entry):
        if sheet_name not in [sheet["name"] for sheet in entry.get("sheets", [])]:
            raise KeyError
        entry["active_sheet"] = sheet_name
        return entry

    try:
        updated_entry = update_import_record(import_id, updater)
    except KeyError:
        return jsonify({"error": "Unknown sheet selected"}), 400
    if not updated_entry:
        return jsonify({"error": "Import session not found"}), 404
    return jsonify(build_import_response(updated_entry))


@app.route("/save_mapping_preset/<table_type>", methods=["POST"])
def save_mapping_preset(table_type):
    if table_type not in ("A", "B"):
        return jsonify({"error": "Invalid table type"}), 400

    payload = request.get_json(silent=True) or request.form
    name = clean_text(payload.get("name"))
    mapping = payload.get("mapping", {})
    if isinstance(mapping, str):
        try:
            mapping = json.loads(mapping)
        except json.JSONDecodeError:
            mapping = {}

    if not name:
        return jsonify({"error": "Preset name is required"}), 400
    preset = save_preset(table_type, name, mapping)
    return jsonify({"status": "success", "preset": preset})


@app.route("/apply_import/<import_id>", methods=["POST"])
def apply_import(import_id):
    payload = request.get_json(silent=True) or {}
    mapping = payload.get("mapping", {})
    preset_name = clean_text(payload.get("preset_name"))
    should_save_preset = bool(payload.get("save_preset"))

    active_division = get_active_division()
    imports = load_json("I", [], division=active_division)
    import_entry = next(
        (entry for entry in imports if entry.get("id") == import_id), None
    )
    if not import_entry:
        return jsonify({"error": "Import session not found"}), 404

    active_sheet = next(
        (
            sheet
            for sheet in import_entry.get("sheets", [])
            if sheet.get("name") == import_entry.get("active_sheet")
        ),
        None,
    )
    if not active_sheet:
        return jsonify({"error": "No active worksheet selected"}), 400

    headers = active_sheet.get("headers", [])
    if not headers:
        return jsonify({"error": "The selected sheet is empty"}), 400

    errors = validate_mapping(headers, mapping)
    if errors:
        return jsonify({"error": " ".join(errors), "validation_errors": errors}), 400

    table_type = import_entry["table_type"]
    existing_records = read_transactions(table_type, active_division)
    known_ids = get_all_known_ids()

    imported = 0
    skipped = 0
    duplicates = 0
    validation_errors = []

    for index, row in enumerate(active_sheet.get("rows", []), start=2):
        transaction, row_error = normalize_row(
            row=row,
            table_type=table_type,
            mapping=mapping,
            import_entry=import_entry,
            active_sheet_name=active_sheet["name"],
        )
        if row_error:
            skipped += 1
            validation_errors.append(f"Row {index}: {row_error}")
            continue
        if transaction["id"] in known_ids:
            duplicates += 1
            continue
        existing_records.append(transaction)
        known_ids.add(transaction["id"])
        imported += 1

    save_json(table_type, existing_records, division=active_division)

    if should_save_preset and preset_name:
        save_preset(table_type, preset_name, mapping)

    return jsonify(
        {
            "status": "success",
            "summary": {
                "rows_read": len(active_sheet.get("rows", [])),
                "rows_imported": imported,
                "rows_skipped": skipped,
                "duplicates": duplicates,
                "validation_errors": validation_errors[:10],
            },
        }
    )


@app.route("/add_to_queue", methods=["POST"])
def add_to_queue():
    selected_ids = request.form.getlist("transaction_items")
    active_division = get_active_division()
    table_a = read_transactions("A", active_division)
    table_b = read_transactions("B", active_division)
    queue = read_queue(active_division)

    for selected_id in selected_ids:
        prefix, actual_id = selected_id.split("_", 1)
        target_table = table_a if prefix == "A" else table_b

        item = next((x for x in target_table if str(x["id"]) == actual_id), None)
        if item:
            queue.append(item)
            target_table.remove(item)

    save_json("A", table_a, division=active_division)
    save_json("B", table_b, division=active_division)
    save_json("Q", queue, division=active_division)
    return redirect(url_for("index"))


@app.route("/remove_from_queue/<source>/<item_id>", methods=["POST"])
def remove_from_queue(source, item_id):
    if source not in ("A", "B"):
        return redirect(url_for("index"))

    active_division = get_active_division()
    queue = read_queue(active_division)
    table_data = read_transactions(source, active_division)
    item = next(
        (x for x in queue if str(x["id"]) == item_id and x["source_type"] == source),
        None,
    )
    if item:
        queue.remove(item)
        table_data.append(item)

    save_json("Q", queue, division=active_division)
    save_json(source, table_data, division=active_division)
    return redirect(url_for("index"))


@app.route("/clear_queue", methods=["POST"])
def process_group():
    active_division = get_active_division()
    queue = read_queue(active_division)
    if not queue:
        return redirect(url_for("index"))

    processed = read_processed(active_division)
    items_a = [i for i in queue if i["source_type"] == "A"]
    items_b = [i for i in queue if i["source_type"] == "B"]
    group_sum = calculate_variance(queue)

    new_group = {
        "group_id": str(uuid.uuid4())[:6].upper(),
        "timestamp": utc_now(),
        "items_a": items_a,
        "items_b": items_b,
        "group_sum": f"{group_sum:.2f}",
    }

    processed.insert(0, new_group)
    save_json("P", processed, division=active_division)
    save_json("Q", [], division=active_division)
    return redirect(url_for("index"))


@app.route("/download_report")
def download_report():
    active_division = get_active_division()
    processed = read_processed(active_division)
    now = utc_now()
    total_variance = sum(
        calculate_variance(group.get("items_a", []) + group.get("items_b", []))
        for group in processed
    )

    report = "====================================================\n"
    report += "          STEALTH RECONCILIATION REPORT             \n"
    report += f"          DIVISION: {active_division:<32}\n"
    report += f"          GENERATED: {now}                          \n"
    report += "====================================================\n\n"
    report += "SUMMARY:\n"
    report += f"Total Recos: {len(processed)}\n"
    report += f"Net Variance: {total_variance:.2f}\n"
    report += "-" * 52 + "\n\n"

    for group in processed:
        group_variance = calculate_variance(
            group.get("items_a", []) + group.get("items_b", [])
        )
        report += f"RECO: {group['group_id']} | DIFF: {group_variance:.2f}\n"
        for item in group["items_a"]:
            report += (
                f" [GL]   {item['description'][:25]:<25} {item['amount']:>10.2f}\n"
            )
        for item in group["items_b"]:
            report += (
                f" [BANK] {item['description'][:25]:<25} {item['amount']:>10.2f}\n"
            )
        report += "-" * 30 + "\n"

    return Response(
        report,
        mimetype="text/plain",
        headers={
            "Content-disposition": f"attachment; filename=Recon_Report_{active_division}.txt"
        },
    )


@app.route("/download_report_xlsx")
def download_report_xlsx():
    active_division = get_active_division()
    processed = read_processed(active_division)
    workbook = build_processed_export_workbook(processed)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = (
        f"Recon_Report_{active_division}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    )
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-disposition": f"attachment; filename={filename}"},
    )


@app.route("/export_transactions/<source_type>")
def export_transactions(source_type):
    if source_type not in ("A", "B"):
        return redirect(url_for("index"))

    active_division = get_active_division()
    transactions = read_transactions(source_type, active_division)
    workbook = build_transactions_export_workbook(source_type, transactions)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    label = "GL" if source_type == "A" else "Bank"
    filename = f"{active_division}_{label}_Transactions_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-disposition": f"attachment; filename={filename}"},
    )


@app.route("/add_adjustment", methods=["POST"])
def add_adjustment():
    item_name = clean_text(request.form.get("item"))
    amount = parse_amount(request.form.get("amount"))
    if not item_name or amount is None:
        return jsonify({"error": "Adjustment note and amount are required"}), 400

    active_division = get_active_division()
    queue = read_queue(active_division)
    source_type = "A" if amount > 0 else "B"
    queue.append(
        build_transaction(
            source_type=source_type,
            import_batch_id="manual-adjustment",
            file_name="manual",
            sheet_name="manual",
            mapped_fields={"description": "manual", "amount": "manual"},
            raw_row={"item": item_name, "amount": amount},
            description=f"ADJ: {item_name}",
            amount=amount,
            notes="Manual balancing entry",
            existing_id=str(uuid.uuid4())[:8],
        )
    )
    save_json("Q", queue, division=active_division)
    return jsonify({"status": "success"})


@app.route("/remove_group/<group_id>", methods=["POST"])
def remove_group(group_id):
    active_division = get_active_division()
    processed = read_processed(active_division)
    updated_processed = [group for group in processed if group["group_id"] != group_id]
    save_json("P", updated_processed, division=active_division)
    return redirect(url_for("index"))


init_db()


if __name__ == "__main__":
    app.run(debug=True, port=5001)
